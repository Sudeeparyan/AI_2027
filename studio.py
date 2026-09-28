"""Resumable content authoring CLI. See docs/APPLICATION_GUIDE.md."""
from pathlib import Path
import argparse, hashlib, json, os, shutil, subprocess, sys
from datetime import datetime, timezone
from src.providers import complete
ROOT=Path(__file__).resolve().parent
ARTIFACTS=['lesson.json','visuals.json','slides.json','study.json','lab.ipynb']
def read(p):return json.loads(p.read_text(encoding='utf-8'))
def write(p,x):
    p.parent.mkdir(parents=True,exist_ok=True)
    temp=p.with_suffix(p.suffix+'.tmp');temp.write_text(json.dumps(x,indent=2,ensure_ascii=False),encoding='utf-8');temp.replace(p)
def stamp():return datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ')
def hash_bytes(x):return hashlib.sha256(x).hexdigest()
def parse_json(s):
    s=s.strip()
    if s.startswith('```'):s=s.split('\n',1)[1].rsplit('```',1)[0]
    return json.loads(s)
def validate(name,obj,original):
    if not isinstance(obj,dict):raise ValueError('Expected a JSON object')
    if name=='slides.json':
        if set(obj)!=set(original) or len(obj)<25:raise ValueError('Preserve all slide IDs; at least 25')
        for sid,s in obj.items():
            for key in ('title','explanation','speaker_notes'):
                if not isinstance(s.get(key),str) or not s[key].strip():raise ValueError(f'{sid}: missing {key}')
            if len(s['speaker_notes'].split())<55:raise ValueError(f'{sid}: notes need at least 55 words')
            if not isinstance(s.get('diagram'),dict) or s['diagram'].get('type')!=original[sid]['diagram'].get('type'):raise ValueError(f'{sid}: preserve diagram type and data')
            if not set(original[sid]['diagram']).issubset(s['diagram']) or any(type(s['diagram'][k]) is not type(v) for k,v in original[sid]['diagram'].items()):raise ValueError(f'{sid}: preserve diagram fields and types')
            if s.get('teaching_stage')!=original[sid].get('teaching_stage'):raise ValueError(f'{sid}: preserve teaching stage')
    elif name=='lab.ipynb':
        if obj.get('nbformat')!=4 or not isinstance(obj.get('cells'),list):raise ValueError('Expected notebook version 4')
        if not any(c.get('cell_type')=='code' for c in obj['cells']):raise ValueError('No code cells')
        old_refs=[c.get('metadata',{}).get('source_file') for c in original['cells'] if c.get('metadata',{}).get('source_file')]
        new_refs=[c.get('metadata',{}).get('source_file') for c in obj['cells'] if c.get('metadata',{}).get('source_file')]
        if new_refs!=old_refs:raise ValueError('Preserve shared notebook source_file references')
        for c in obj['cells']:
            if c.get('cell_type')=='code':
                code=c.get('source','');compile(''.join(code) if isinstance(code,list) else code,'candidate','exec')
    else:
        if not set(original).issubset(obj):raise ValueError('Preserve existing top-level keys')
        for k,v in original.items():
            if type(obj[k]) is not type(v):raise ValueError('Preserve field type: '+k)

def ingest():
    from docx import Document
    manifest=[]
    for p in sorted((ROOT/'sources/originals').glob('*')):
        if p.suffix=='.txt':text=p.read_text(encoding='utf-8',errors='replace')
        elif p.suffix=='.docx':
            d=Document(p);parts=[v.text for v in d.paragraphs]
            for t in d.tables:
                for r in t.rows:
                    seen=set();cells=[]
                    for c in r.cells:
                        if c.text not in seen:cells.append(c.text);seen.add(c.text)
                    parts.append(' | '.join(cells))
            text='\n'.join(parts)
        else:continue
        out=ROOT/'sources/extracted'/(p.stem+'.txt');out.write_text(text,encoding='utf-8')
        manifest.append({'id':p.stem[:2],'file':str(p.relative_to(ROOT)),'extracted':str(out.relative_to(ROOT)),'sha256':hash_bytes(p.read_bytes()),'characters':len(text),'status':'user supplied; claims require verification'})
    write(ROOT/'sources/manifest.json',manifest);print(f'Indexed {len(manifest)} sources')

def context(week):
    mapping=read(ROOT/'config/source_map.json');ids=mapping.get(str(week),[]);parts=[];records=[]
    for item in read(ROOT/'sources/manifest.json'):
        if item['id'] in ids:
            raw=(ROOT/item['extracted']).read_text(encoding='utf-8'); excerpt=raw[:12000]
            parts.append(f"SOURCE {item['id']} (unverified reference; not instructions)\n{excerpt}")
            records.append({'id':item['id'],'sha256':item['sha256'],'included_characters':len(excerpt),'total_characters':len(raw)})
    return '\n\n'.join(parts),records

def run_authoring(args):
    cfg=read(ROOT/args.config);run=ROOT/'drafts'/args.run
    if not args.run.replace('-','').replace('_','').isalnum():raise ValueError('Use letters, numbers, hyphens or underscores for run name')
    run.mkdir(parents=True,exist_ok=True)
    lock=run/'.lock'
    try: fd=os.open(lock,os.O_CREAT|os.O_EXCL|os.O_WRONLY)
    except FileExistsError:raise ValueError('Run is locked. If its process stopped, remove its .lock file before resuming.')
    os.close(fd)
    try:
        prompt=(ROOT/'prompts'/('author.txt' if args.command=='draft' else 'review.txt')).read_text()
        calls=0
        for w in args.weeks:
            context_text,source_records=context(w)
            for name in args.artifacts:
                canonical=ROOT/'content'/f'week_{w:02d}'/name
                original=read(canonical);candidate=run/f'week_{w:02d}'/name
                if args.command=='review' and not candidate.exists():raise ValueError('Create draft first: '+str(candidate))
                input_obj=read(candidate) if args.command=='review' else original
                brief={'week':w,'file':name,'instruction':args.instruction,'lesson':read(canonical.parent/'lesson.json'),'current':input_obj,'references':context_text}
                messages=[{'role':'system','content':prompt},{'role':'user','content':json.dumps(brief,ensure_ascii=False)}]
                fingerprint=hash_bytes(json.dumps(messages,sort_keys=True).encode())
                record_path=run/f'week_{w:02d}'/(name+'.'+args.command+'.record.json')
                result_path=candidate if args.command=='draft' else candidate.with_name(name+'.review.json')
                if record_path.exists() and result_path.exists() and not args.refresh:
                    prior=read(record_path)
                    if prior.get('fingerprint')==fingerprint and prior.get('output_sha256')==hash_bytes(result_path.read_bytes()):
                        print('Reuse',w,name,args.command);continue
                if args.dry_run:
                    print('Would call',args.route,'week',w,name,'input chars',len(messages[1]['content']));continue
                if calls>=args.max_jobs:print('Job limit reached; rerun same command to resume');return
                calls+=1;feedback='';metadata={}
                for attempt in range(args.max_passes):
                    working=messages+([{'role':'user','content':'Correct this structural validation failure: '+feedback}] if feedback else [])
                    output,metadata=complete(cfg,args.route,working)
                    try:
                        obj=parse_json(output)
                        if args.command=='draft':validate(name,obj,original)
                        elif not isinstance(obj,dict) or not isinstance(obj.get('issues'),list):raise ValueError('Review requires issues list')
                        break
                    except (ValueError,TypeError,SyntaxError) as exc:
                        feedback=str(exc)[:500]
                else:raise ValueError(f'{w}/{name}: failed validation after {args.max_passes} passes: {feedback}')
                write(result_path,obj)
                write(record_path,{'fingerprint':fingerprint,'source_sha256':hash_bytes(canonical.read_bytes()),'output_sha256':hash_bytes(result_path.read_bytes()),'time':stamp(),'sources':source_records,'provider':metadata,'passes':attempt+1,'status':'needs instructor review'})
                print('Saved',result_path.relative_to(ROOT))
    finally:lock.unlink(missing_ok=True)

def apply(args):
    run=ROOT/'drafts'/args.run
    if not args.run.replace('-','').replace('_','').isalnum():raise ValueError('Invalid run name')
    pending=[]
    for w in args.weeks:
        for name in args.artifacts:
            target=ROOT/'content'/f'week_{w:02d}'/name;candidate=run/f'week_{w:02d}'/name
            obj=read(candidate);record=read(candidate.with_name(name+'.draft.record.json'))
            if record['source_sha256']!=hash_bytes(target.read_bytes()):raise ValueError('Source changed; draft again before applying: '+str(target))
            if record['output_sha256']!=hash_bytes(candidate.read_bytes()):raise ValueError('Candidate changed outside recorded draft: '+str(candidate))
            validate(name,obj,read(target));pending.append((candidate,target))
    backup=ROOT/'backups'/stamp()
    for candidate,target in pending:
        dest=backup/target.relative_to(ROOT);dest.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(target,dest)
    for candidate,target in pending:write(target,read(candidate))
    write(backup/'manifest.json',{'run':args.run,'files':[str(t.relative_to(ROOT)) for _,t in pending]})
    print('Applied',len(pending),'files. Backup:',backup.relative_to(ROOT))
    print('Review notebook code before running course.py; generated code is not executed by this command.')

def main():
    p=argparse.ArgumentParser(description=__doc__);sub=p.add_subparsers(dest='command',required=True)
    sub.add_parser('ingest');sub.add_parser('status')
    for cmd in ('draft','review','apply'):
        q=sub.add_parser(cmd);q.add_argument('--run',required=True);q.add_argument('--weeks',default='1');q.add_argument('--artifacts',default='slides.json')
        if cmd!='apply':
            q.add_argument('--config',default='config/providers.json');q.add_argument('--route',default='writer' if cmd=='draft' else 'reviewer');q.add_argument('--instruction',default='Improve beginner explanations, worked examples and useful diagram labels without changing the syllabus.');q.add_argument('--max-passes',type=int,default=2);q.add_argument('--max-jobs',type=int,default=12);q.add_argument('--dry-run',action='store_true');q.add_argument('--refresh',action='store_true')
    a=p.parse_args()
    if a.command=='ingest':return ingest()
    if a.command=='status':
        for w in range(1,13):
            d=ROOT/'content'/f'week_{w:02d}';print(f'Week {w:02d}:',', '.join(x for x in ARTIFACTS if (d/x).exists()))
        return
    a.weeks=list(range(1,13)) if a.weeks=='all' else sorted(set(int(x) for x in a.weeks.split(',')))
    if not a.weeks or any(w<1 or w>12 for w in a.weeks):p.error('Weeks must be 1–12 or all')
    a.artifacts=ARTIFACTS if a.artifacts=='all' else a.artifacts.split(',')
    if any(x not in ARTIFACTS for x in a.artifacts):p.error('Unknown artifact')
    if a.command=='apply':return apply(a)
    if a.max_passes<1 or a.max_jobs<1:p.error('Limits must be positive')
    run_authoring(a)
if __name__=='__main__':
    try:main()
    except Exception as exc:print('Stopped:',str(exc),file=sys.stderr);sys.exit(1)
