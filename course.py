"""Build and review reusable teaching materials. Run `python course.py --help`."""
from pathlib import Path
import argparse, contextlib, csv, hashlib, io, json, shutil, subprocess, sys, tempfile, time, html
from datetime import datetime, timezone
import fitz
from pptx import Presentation
from pptx.util import Inches, Pt
from src import slides, handouts, engaging_slides

ROOT=Path(__file__).resolve().parent
def read(p): return json.loads(p.read_text(encoding='utf-8'))
def write(p,data):
    p.parent.mkdir(parents=True,exist_ok=True)
    p.write_text(json.dumps(data,indent=2,ensure_ascii=False),encoding='utf-8')
def digest(data): return hashlib.sha256(data).hexdigest()
def weeks_arg(value):
    result=list(range(1,13)) if value=='all' else [int(v) for v in value.split(',')]
    if not result or any(w<1 or w>12 for w in result): raise ValueError('Weeks must be 1 to 12, or all.')
    return sorted(set(result))
def fingerprint(week):
    files=sorted((ROOT/'content'/f'week_{week:02d}').glob('*'))
    files+=sorted((ROOT/'src').glob('*.py'))+[ROOT/'course.py',ROOT/'content/references.json']
    if week in (7,11,12):files+=sorted((ROOT/'content/shared').glob('*.py'))
    if week==1: files+=sorted((ROOT/'assets/illustrations').glob('*'))
    design=ROOT/'content/teaching_design/weeks.json'
    weekly_design=json.dumps(read(design)[str(week)],sort_keys=True,ensure_ascii=False).encode() if design.exists() else b''
    return digest(b''.join(p.read_bytes() for p in files if p.is_file())+weekly_design)

def audit(deck,notebook):
    """Objective checks only. Pedagogy and diagram meaning need human review."""
    p=Presentation(deck);issues=[]
    if len(p.slides)<25: issues.append({'kind':'count','message':'Fewer than 25 slides'})
    for i,s in enumerate(p.slides,1):
        if len(s.notes_slide.notes_text_frame.text.split())<55:
            issues.append({'kind':'notes','slide':i,'message':'Speaker explanation is too short'})
        for sh in s.shapes:
            if sh.left<0 or sh.top<0 or sh.left+sh.width>p.slide_width+2500 or sh.top+sh.height>p.slide_height+2500:
                issues.append({'kind':'bounds','slide':i,'shape':sh.shape_id,'message':'Object outside slide'})
            if sh.has_text_frame:
                # Detect long unbroken words likely to split inside a narrow box.
                for para in sh.text_frame.paragraphs:
                    size=para.font.size.pt if para.font.size else 17
                    capacity=max(1,(sh.width/914400*72-15)/(size*.5))
                    if any(len(word)>capacity+1 for word in para.text.split()):
                        issues.append({'kind':'word_wrap','slide':i,'shape':sh.shape_id,'message':'Long word in narrow box'})
    try:
        nb=read(notebook)
        program='\n'.join(''.join(c['source']) if isinstance(c['source'],list) else c['source'] for c in nb['cells'] if c['cell_type']=='code')
        with tempfile.TemporaryDirectory() as temp:
            result=subprocess.run([sys.executable,'-c',program],cwd=temp,capture_output=True,text=True,timeout=30)
        if result.returncode: issues.append({'kind':'notebook','message':result.stderr[-2000:]})
    except Exception as exc: issues.append({'kind':'notebook','message':str(exc)})
    return issues

def repair_layout(deck,issues):
    """Conservative font/position repairs; never change teaching claims."""
    p=Presentation(deck);changes=[]
    for issue in issues:
        if issue['kind'] not in ('word_wrap','bounds'): continue
        s=p.slides[issue['slide']-1]
        sh=next(x for x in s.shapes if x.shape_id==issue['shape'])
        if issue['kind']=='bounds':
            sh.left=max(0,min(sh.left,p.slide_width-sh.width));sh.top=max(0,min(sh.top,p.slide_height-sh.height))
            changes.append({**issue,'fix':'Moved object within slide bounds'})
        else:
            for para in sh.text_frame.paragraphs:
                size=para.font.size.pt if para.font.size else 17
                if size>12: para.font.size=Pt(max(12,size-1))
            changes.append({**issue,'fix':'Reduced diagram label font by 1 pt, minimum 12 pt'})
    p.save(deck);return changes

def export_visuals(deck,out,week,run):
    exe=shutil.which('soffice') or shutil.which('libreoffice')
    if not exe: raise RuntimeError('Install LibreOffice and make soffice available on PATH to render diagrams.')
    with tempfile.TemporaryDirectory(prefix='ai2027_lo_') as profile:
        handout=out/f'Week_{week:02d}_Reading_and_Teaching_Notes.docx'
        subprocess.run([exe,f'-env:UserInstallation={Path(profile).as_uri()}','--headless','--convert-to','pdf','--outdir',str(out),str(deck),str(handout)],check=True,capture_output=True,timeout=180)
    pdf=out/(deck.stem+'.pdf')
    if not pdf.exists(): raise RuntimeError('LibreOffice did not produce a PDF.')
    document=fitz.open(pdf);entries=[];assets=ROOT/'assets/diagrams'/f'week_{week:02d}';assets.mkdir(parents=True,exist_ok=True)
    previews=out/'previews';previews.mkdir(exist_ok=True)
    for idx,page in enumerate(document,1):
        crop=fitz.Rect(5.12*72,1.42*72,12.62*72,6.7*72)
        pix=page.get_pixmap(matrix=fitz.Matrix(1.7,1.7),clip=crop)
        raw=pix.tobytes('png');h=digest(raw)[:12];stem=f's{idx:02d}-{h}'
        png=assets/f'{stem}.png';svg=assets/f'{stem}.svg'
        if not png.exists(): png.write_bytes(raw)
        old=page.cropbox;page.set_cropbox(crop)
        if not svg.exists(): svg.write_text(page.get_svg_image(),encoding='utf-8')
        page.set_cropbox(old)
        page.get_pixmap(matrix=fitz.Matrix(.7,.7)).save(previews/f's{idx:02d}.png')
        entries.append({'slide_id':f's{idx:02d}','png':str(png.relative_to(ROOT)),'svg':str(svg.relative_to(ROOT)),'sha256':digest(raw),'source_run':run.name})
    document.close()
    handout_pdf=handout.with_suffix('.pdf')
    if handout_pdf.exists():
        hd=fitz.open(handout_pdf);hp=out/'handout_previews';hp.mkdir(exist_ok=True)
        for i,page in enumerate(hd,1):page.get_pixmap(matrix=fitz.Matrix(1,1)).save(hp/f'page-{i:02d}.png')
        hd.close()
    # Reuse those exact exported images in the classroom deck. The editable
    # version remains beside it, so future edits do not depend on bitmap editing.
    p=Presentation(deck)
    for s,entry in zip(p.slides,entries):
        for sh in list(s.shapes):
            if sh.left>=Inches(5.1):sh._element.getparent().remove(sh._element)
        s.shapes.add_picture(str(ROOT/entry['png']),Inches(5.12),Inches(1.42),width=Inches(7.5),height=Inches(5.28))
    p.save(out/f'Week_{week:02d}_Slides.pptx')
    write(out/'asset_manifest.json',entries)
    return entries

def build(weeks,max_passes=1,force=False):
    stamp=datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ')
    run=ROOT/'outputs/runs'/stamp;run.mkdir(parents=True)
    latest_path=ROOT/'outputs/latest.json';latest=read(latest_path) if latest_path.exists() else {}
    report={'run':stamp,'weeks':{},'human_review':'pending; automatic checks do not grade educational quality'}
    handouts.SOURCES=read(ROOT/'content/references.json');slides.ASSET_ROOT=ROOT/'assets'
    for n in weeks:
        key=f'week_{n:02d}';content=ROOT/'content'/key;fp=fingerprint(n)
        if not force and latest.get(key,{}).get('fingerprint')==fp:
            report['weeks'][key]={'status':'unchanged; reused current build'};print(key,'unchanged',flush=True);continue
        out=run/key;out.mkdir();w=read(content/'lesson.json')
        if (content/'study.json').exists():w['study']=read(content/'study.json')
        slides.VISUALS[n]=read(content/'visuals.json')
        slides.OVERRIDES=read(content/'slides.json') if (content/'slides.json').exists() else {}
        design=ROOT/'content/teaching_design/weeks.json'
        if design.exists():
            w['design']=read(design)[str(n)]
            engaging_slides.create_deck(w,out,slides.OVERRIDES)
        else:slides.create_deck(w,out)
        native=out/f'Week_{n:02d}_Editable.pptx';(out/f'Week_{n:02d}_Slides.pptx').rename(native)
        if not (content/'slides.json').exists():
            p=Presentation(native);overrides={}
            for i,s in enumerate(p.slides,1):
                overrides[f's{i:02d}']={'title':s.shapes[0].text,'explanation':s.shapes[2].text,'speaker_notes':s.notes_slide.notes_text_frame.text}
            write(content/'slides.json',overrides)
        handouts.make_doc(w,out)
        lab=out/f'Week_{n:02d}_Lab.ipynb';nb=read(content/'lab.ipynb')
        for cell in nb['cells']:
            if cell.get('metadata',{}).get('source_file'):
                reference=(ROOT/cell['metadata']['source_file']).resolve()
                if not reference.is_relative_to((ROOT/'content/shared').resolve()):
                    raise ValueError('Notebook source_file must refer to content/shared/')
                cell['source']=reference.read_text(encoding='utf-8')
        if w.get('design'):
            de=w['design']
            intro=(f'# Week {n:02d}: {w["title"]} — the case\n\n'
                   f'**Start here:** {de["hook"]}\n\n'
                   f'**Your goal:** {w["goal"]}\n\n'
                   'Run the prepared cells from top to bottom. Predict the result before changing one input. '
                   'These small demonstrations isolate an idea; they are not trained production models.\n\n'
                   '**Concepts to point to:** '+', '.join(t[0] for t in w['terms'])+'.')
            reflection=('## Explain and transfer\n\n'
                        f'1. What did you change, and what happened? {w["change"]}\n'
                        f'2. Which diagram block produced the difference?\n'
                        f'3. What is one limitation of the demonstration?\n'
                        f'4. Fresh situation: {de["transfer"]}\n\n'
                        f'**Evidence to save:** {w["output"]}')
            nb['cells'].insert(0,{'cell_type':'markdown','id':f'w{n:02d}-story','metadata':{},'source':intro})
            nb['cells'].append({'cell_type':'markdown','id':f'w{n:02d}-transfer','metadata':{},'source':reflection})
        write(lab,nb)
        passes=[]
        for attempt in range(1,max_passes+1):
            issues=audit(native,lab);record={'pass':attempt,'issues':issues,'repairs':[]};passes.append(record)
            if not issues: break
            if attempt<max_passes:record['repairs']=repair_layout(native,issues)
        if passes[-1]['issues']:
            report['weeks'][key]={'passes':passes,'remaining_issues':len(passes[-1]['issues']),'status':'failed; current release preserved'}
            write(run/'report.json',report)
            raise RuntimeError(f'{key} still has failed checks; current release was not replaced. See {run / "report.json"}')
        entries=export_visuals(native,out,n,run)
        report['weeks'][key]={'passes':passes,'remaining_issues':len(passes[-1]['issues']),'assets':len(entries),'status':'review needed' if passes[-1]['issues'] else 'automatic checks passed'}
        write(out/'build_report.json',{'run':stamp,**report['weeks'][key]})
        snap=run/'source_snapshot'/key;shutil.copytree(content,snap)
        current=ROOT/'outputs/current'/key
        if current.exists():shutil.rmtree(current)
        shutil.copytree(out,current)
        latest[key]={'fingerprint':fingerprint(n),'run':stamp,'output':str(current.relative_to(ROOT))}
        write(latest_path,latest);write(run/'report.json',report)
        print(key,report['weeks'][key]['status'],len(entries),'diagram images',flush=True)
    shutil.copytree(ROOT/'src',run/'source_snapshot/src',ignore=shutil.ignore_patterns('__pycache__'))
    shutil.copy2(ROOT/'course.py',run/'source_snapshot/course.py')
    shutil.copy2(ROOT/'content/references.json',run/'source_snapshot/references.json')
    if (ROOT/'content/teaching_design').exists():
        shutil.copytree(ROOT/'content/teaching_design',run/'source_snapshot/teaching_design',ignore=shutil.ignore_patterns('__pycache__'))
    if (ROOT/'content/shared').exists():shutil.copytree(ROOT/'content/shared',run/'source_snapshot/shared',ignore=shutil.ignore_patterns('__pycache__'))
    write(run/'report.json',report);write(latest_path,latest)
    gallery=['<!doctype html><meta charset="utf-8"><title>AI2027 review gallery</title><style>body{font:17px system-ui;max-width:1400px;margin:36px auto;background:#f4f7fa;color:#182b40}section{margin:40px 0}img{width:100%;border:1px solid #ddd}article{display:inline-block;vertical-align:top;width:31%;margin:1%}a{color:#255e9c}</style><h1>AI2027: current slide review</h1><p>Use the CSV in reviews/ to record teaching and visual feedback. A successful build does not replace that review.</p>']
    catalog=[]
    for key in sorted(latest):
        folder=ROOT/latest[key]['output'];week=int(key[-2:])
        gallery.append(f'<section><h2>Week {week:02d}</h2><p><a href="{key}/Week_{week:02d}_Slides.pptx">Classroom deck</a> · <a href="{key}/Week_{week:02d}_Editable.pptx">Editable deck</a></p>')
        for preview in sorted((folder/'previews').glob('*.png')):
            gallery.append(f'<article><img loading="lazy" src="{key}/previews/{preview.name}" alt="Week {week} slide {preview.stem}"><p>{preview.stem}</p></article>')
        gallery.append('</section>')
        catalog+=read(folder/'asset_manifest.json')
    (ROOT/'outputs/current/index.html').write_text(''.join(gallery),encoding='utf-8')
    write(ROOT/'assets/catalog.json',catalog)
    review=ROOT/'reviews'/f'{stamp}.csv'
    with review.open('w',newline='',encoding='utf-8') as f:
        writer=csv.writer(f);writer.writerow(['week','slide','clear_to_beginner','diagram_correct','example_helpful','notes_teachable','issue','proposed_change','status'])
        for n in weeks:
            for i in range(1,len(read(ROOT/'content'/f'week_{n:02d}'/'slides.json'))+1):writer.writerow([n,f's{i:02d}','','','','','','','pending'])
    print('Report:',run/'report.json');print('Human review:',review)
    return report

if __name__=='__main__':
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('command',choices=['build','improve','watch'])
    ap.add_argument('--weeks',default='all',help='all or comma-separated numbers, e.g. 1,2')
    ap.add_argument('--max-passes',type=int,default=3,help='Bounded layout-repair passes, 1 to 5')
    ap.add_argument('--force',action='store_true',help='Rebuild unchanged inputs')
    args=ap.parse_args()
    if not 1<=args.max_passes<=5:ap.error('--max-passes must be between 1 and 5')
    selected=weeks_arg(args.weeks)
    if args.command=='watch':
        print('Watching content and Python files. Save an edit to rebuild. Ctrl+C stops.')
        previous={}
        try:
            while True:
                current={n:fingerprint(n) for n in selected}
                changed=[n for n in selected if previous.get(n)!=current[n]]
                if changed:
                    build(changed,args.max_passes,args.force)
                    previous={n:fingerprint(n) for n in selected}
                time.sleep(2)
        except KeyboardInterrupt:print('Stopped. All completed runs remain saved.')
    else:build(selected,args.max_passes if args.command=='improve' else 1,args.force)
