"""Validate the 12-week release; run after course.py improve --weeks all."""
from pathlib import Path
import hashlib,json,sys
from pptx import Presentation
from docx import Document
import fitz
ROOT=Path(__file__).resolve().parent

def verify():
    errors=[];weeks=[]
    for w in range(1,13):
        key=f'week_{w:02d}';folder=ROOT/'outputs/current'/key
        try:
            slides=Presentation(folder/f'Week_{w:02d}_Slides.pptx')
            native=Presentation(folder/f'Week_{w:02d}_Editable.pptx')
            source=json.loads((ROOT/'content'/key/'slides.json').read_text())
            assert len(slides.slides)==len(native.slides)==len(source)==40,'slide count'
            assets=json.loads((folder/'asset_manifest.json').read_text());assert len(assets)==40,'asset count'
            for i,(s,entry) in enumerate(zip(slides.slides,assets),1):
                assert len(s.notes_slide.notes_text_frame.text.split())>=55,f'slide {i} notes'
                assert source[f's{i:02d}']['title'] in '\n'.join(x.text for x in s.shapes if x.has_text_frame),f'slide {i} title'
                pictures=[x for x in s.shapes if x.shape_type==13];assert len(pictures)==1,f'slide {i} picture'
                raw=(ROOT/entry['png']).read_bytes();assert hashlib.sha256(raw).hexdigest()==entry['sha256'],f'slide {i} asset hash'
                assert pictures[0].image.blob==raw,f'slide {i} embedded asset'
                assert (ROOT/entry['svg']).is_file(),f'slide {i} svg'
            doc=Document(folder/f'Week_{w:02d}_Reading_and_Teaching_Notes.docx')
            text='\n'.join(p.text for p in doc.paragraphs)
            assert 'Model answers and marking guidance' in text,'exam answers'
            assert 'Official reading and video support' in text,'resources'
            assert 'Teach the week as a story' in text,'teaching story'
            assert all(term[0] in text for term in json.loads((ROOT/'content'/key/'lesson.json').read_text())['terms']),'concept explanations'
            nb=json.loads((folder/f'Week_{w:02d}_Lab.ipynb').read_text())
            ids=[c['id'] for c in nb['cells']];assert len(ids)==len(set(ids)),'duplicate cell IDs'
            assert nb['nbformat']==4 and any(c['cell_type']=='code' for c in nb['cells']),'notebook structure'
            build=json.loads((folder/'build_report.json').read_text());assert build['remaining_issues']==0,'build issues'
            page_counts={}
            for pdf in folder.glob('*.pdf'):
                with fitz.open(pdf) as rendered:
                    assert len(rendered)>0,'empty PDF'
                    for pn,page in enumerate(rendered,1):
                        assert page.get_text().strip(),f'blank PDF page {pn}'
                        for word in page.get_text('words'):
                            x0,y0,x1,y1=word[:4]
                            assert x0>=-1 and y0>=-1 and x1<=page.rect.width+1 and y1<=page.rect.height+1,f'out-of-page text {pn}'
                    page_counts[pdf.name]=len(rendered)
            assert page_counts.get(f'Week_{w:02d}_Editable.pdf')==40,'slide PDF pages'
            assert f'Week_{w:02d}_Reading_and_Teaching_Notes.pdf' in page_counts,'handout PDF'
            weeks.append({'week':w,'slides':40,'pdf_pages':page_counts,'checks':'passed'})
        except Exception as exc:errors.append({'week':w,'error':str(exc)})
    result={'weeks':weeks,'errors':errors,'scope':'Structural checks, source/asset consistency and PDF bounds. Does not certify educational accuracy.'}
    out=ROOT/'reviews/FINAL_STRUCTURAL_CHECK.json';out.write_text(json.dumps(result,indent=2))
    print(json.dumps({'weeks_passed':len(weeks),'errors':errors},indent=2));return bool(errors)
if __name__=='__main__':sys.exit(verify())
