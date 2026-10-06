"""Build the revised curriculum from one canonical lesson per week.

Usage: python build_revised.py --weeks all
Requires python-docx. PPTX creation uses the artifact-tool runtime through Node.
"""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
import argparse, hashlib, html, json, re, shutil, subprocess, sys
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
CANON=ROOT/'content/revised'
STAGE=ROOT/'outputs/revised'
BUILD=ROOT/'.build/revised'
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

def read(p): return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def write(p,data):
    p.parent.mkdir(parents=True,exist_ok=True)
    p.write_text(json.dumps(data,indent=2,ensure_ascii=False),encoding='utf-8')
def sentences(text): return re.split(r'(?<=[.!?])\s+',text.strip())
def brief(text,limit=48):
    parts=sentences(text); result=[]; count=0
    for part in parts:
        words=len(part.split())
        if result and count+words>limit: break
        result.append(part);count+=words
    return result

def make_slides(w):
    refs='\n'.join(w['references'])
    def slide(title,body,notes,kind='text',**extra):
        return dict(title=title,body=body,notes=notes+'\n\nFurther reading\n'+refs,kind=kind,**extra)
    s=[slide(w['title'],[f'Week {w["n"]:02d}', 'Generative AI MSc'],w['overview'],'cover'),
       slide('Learning goals',w['objectives'],w['overview']+' '+w['prerequisites'])]
    for c in w['concepts']:
        s.append(slide(c['title'],brief(c['explanation']),c['explanation']))
        s.append(slide(c['title']+' in practice',brief(c['example'],65),c['example']+' '+c['explanation'],'example'))
        s.append(slide('Discuss '+c['title'].lower(),[c['question']], 'Give students two minutes to think and discuss in pairs. Expected answer: '+c['answer']+' '+c['example'],'question'))
        s.append(slide(c['title']+' check',[c['answer']],c['correction']+' '+c['explanation'],'comparison',table=[['Claim to examine','More accurate explanation'],[c['misconception'],c['correction']]]))
    d=w['demo'];s.append(slide(d['title'],d['steps'],d['expected']+' '+d['offline'],'demo'))
    s.append(slide('What the demonstration shows',[d['expected'],d['offline']],d['expected']+' '+d['offline']+' '+w['concepts'][0]['example'],'example'))
    for i,q in enumerate(w['practice'],1):
        s.append(slide(f'Practice question {i}',[q['question']], 'Ask students to write an answer before discussion. Model answer: '+q['answer'],'question'))
    lab=w['lab'];s.append(slide(lab['title'],lab['steps'],lab['goal']+' Deliverable: '+lab['deliverable'],'lab'))
    s.append(slide('Evidence to keep',[lab['deliverable']],lab['goal']+' '+w['extension'],'example'))
    s.append(slide('Reading and independent practice',['Review the four concepts in the notes.','Complete the notebook and explain one failure.','Use the linked official readings and optional video questions.'], 'Optional extension: '+w['extension']+' Sources were reviewed on 28 September 2026. Check current tool availability before teaching.'))
    assert len(s)==26
    for i,row in enumerate(s,1):row['id']=f's{i:02d}'
    return s

def setup_doc(title,subtitle):
    d=Document();sec=d.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.7)
    sec.left_margin=sec.right_margin=Inches(.8)
    for style in ['Normal','Title','Heading 1','Heading 2','List Bullet','List Number']:
        d.styles[style].font.name='Arial'
    d.styles['Normal'].font.size=Pt(11)
    d.styles['Normal'].paragraph_format.space_after=Pt(7)
    d.styles['Normal'].paragraph_format.line_spacing=1.08
    d.styles['Title'].font.size=Pt(25);d.styles['Title'].font.color.rgb=RGBColor(0,0,0)
    for name,size in [('Heading 1',16),('Heading 2',12)]:
        d.styles[name].font.size=Pt(size);d.styles[name].font.color.rgb=RGBColor(24,61,84)
        d.styles[name].paragraph_format.keep_with_next=True
    for border in d.styles.element.xpath('.//w:pBdr'):border.getparent().remove(border)
    d.add_paragraph(title,'Title');d.add_paragraph(subtitle)
    footer=sec.footer.paragraphs[0];footer.alignment=2
    footer.add_run('Generative AI    ')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    return d

def add_link(doc,label,url):
    p=doc.add_paragraph();hyper=OxmlElement('w:hyperlink')
    rid=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    hyper.set(qn('r:id'),rid);run=OxmlElement('w:r');props=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'1E5F88');props.append(color);run.append(props)
    text=OxmlElement('w:t');text.text=label;run.append(text);hyper.append(run);p._p.append(hyper)
    return p

def source_catalog():
    items=read(ROOT/'research/verified_sources_2026.json')['sources']
    other=ROOT/'research/openai_sources_2026.json'
    if other.exists():items+=read(other)['sources']
    return {r['url']:r for r in items}

def make_notes(w,folder):
    d=setup_doc(f'Week {w["n"]:02d} {w["title"]}','Reference notes and teaching guide')
    d.add_paragraph(w['overview'])
    d.add_heading('Learning goals',1)
    for goal in w['objectives']:d.add_paragraph(goal,'List Bullet')
    d.add_paragraph('Learning outcome links: '+', '.join(w['outcomes'])+'. The original descriptor learning outcomes and assessment scheme remain unchanged.')
    d.add_heading('Preparation and class plan',1);d.add_paragraph(w['prerequisites'])
    d.add_paragraph('Lecture 120 minutes: 10 minutes retrieval practice, four 20-minute concept cycles, 20 minutes demonstration and 10 minutes exit discussion. Each concept cycle combines explanation, a worked example and a student question.')
    d.add_paragraph('Lab 120 minutes: 15 minutes setup and prediction, 25 minutes guided demonstration, 45 minutes independent change, 20 minutes comparison and 15 minutes explanation of results. Use the remaining independent study time for the linked readings, unfinished exercises and portfolio evidence.')
    d.add_heading('Concepts and worked examples',1)
    for c in w['concepts']:
        d.add_heading(c['title'],2);d.add_paragraph(c['explanation'])
        d.add_paragraph('Worked example. '+c['example'])
        d.add_paragraph('Common misconception. '+c['misconception'])
        d.add_paragraph('Correction. '+c['correction'])
        d.add_paragraph('Discuss. '+c['question'])
        d.add_paragraph('Answer guide. '+c['answer'])
    d.add_heading('Classroom demonstration',1);demo=w['demo']
    d.add_heading(demo['title'],2)
    for step in demo['steps']:d.add_paragraph(step,'List Number')
    d.add_paragraph('Expected observation. '+demo['expected'])
    d.add_paragraph('If a live service is unavailable. '+demo['offline'])
    d.add_heading('Guided lab and practice',1);lab=w['lab']
    d.add_paragraph(lab['title']+'. '+lab['goal'])
    for step in lab['steps']:d.add_paragraph(step,'List Number')
    d.add_paragraph('Submit as formative evidence. '+lab['deliverable'])
    if w['n'] in (2,8):
        d.add_paragraph('Required model training. Complete the additional guided PyTorch training notebook in this folder. The small offline calculations alone do not demonstrate model implementation and fine-tuning. The saved training evidence supports discussion when setup fails, but is not a substitute for the learner conducting and explaining the required exercise.')
    d.add_paragraph('Record the exact input, tool or model, date, settings, source and result. Label saved responses, synthetic data and illustrative calculations. Live AI-tool responses may vary. Never claim a simulation is an actual model run.')
    d.add_heading('Practice questions',1)
    for i,q in enumerate(w['practice'],1):d.add_paragraph(f'{i}. {q["question"]}')
    d.add_paragraph('Attempt these without the answer guide. These original questions are formative practice, not an approved examination paper.')
    d.add_heading('Answer guide',1)
    for i,q in enumerate(w['practice'],1):d.add_paragraph(f'{i}. {q["answer"]}')
    d.add_heading('Independent extension',1);d.add_paragraph(w['extension'])
    d.add_heading('Reading and video support',1)
    catalog=source_catalog()
    for url in w['references']:
        r=catalog.get(url,{})
        add_link(d,r.get('title',url),url)
        if r:d.add_paragraph(r.get('curriculum_use',r.get('use','Read the sections related to this week.')))
    videos=read(ROOT/'research/verified_sources_2026.json').get('videos',[])
    if videos:
        video=videos[0 if w['n']<5 else min(1,len(videos)-1)]
        add_link(d,'Optional video or video series '+video['title'],video['url'])
        d.add_paragraph('Viewing task: explain one relevant idea, give your own example and identify one limitation. Video metadata or linking pages were checked; full video content was not reviewed. The written references provide the verified reading path.')
    d.add_paragraph('Research review date 28 September 2026. Recheck service availability, interfaces and account requirements before 2027 delivery. The examples and questions in these notes are original classroom material. Sources inform concepts and further reading; no vendor endorsement is implied.')
    d.save(folder/f'Week_{w["n"]:02d}_Reading_and_Teaching_Notes.docx')

def patch_descriptor():
    source=ROOT/'context/Descriptor - Generative AI-old.docx'
    target=STAGE/'Descriptor - Generative AI revised 7.3.docx'
    rows=read(ROOT/'research/section_7_3.json')
    if isinstance(rows,dict):rows=rows.get('rows',rows.get('weeks'))
    with ZipFile(source) as zin:
        original=zin.read('word/document.xml');tree=etree.fromstring(original)
        body=tree.find('w:body',NS);table=body[12]
        assert table.tag==qn('w:tbl')
        oldrows=table.findall('w:tr',NS);assert len(oldrows)==13
        before=[etree.tostring(el) for el in body]
        for row,values in zip(oldrows[1:],rows):
            strings=[values['topic'],str(values['week']),values['detail'],values['tutorial']]
            for cell,value in zip(row.findall('w:tc',NS),strings):
                paras=cell.findall('w:p',NS)
                p=paras[0];props=p.find('w:pPr',NS)
                runs=p.findall('w:r',NS);rpr=runs[0].find('w:rPr',NS) if runs else None
                for child in list(p):
                    if child is not props:p.remove(child)
                run=etree.SubElement(p,qn('w:r'))
                if rpr is not None:run.append(deepcopy(rpr))
                t=etree.SubElement(run,qn('w:t'));t.text=value
                for extra in paras[1:]:cell.remove(extra)
        for i,el in enumerate(body):
            if i!=12:assert etree.tostring(el)==before[i],f'Unexpected change outside 7.3: {i}'
        edited=etree.tostring(tree,encoding='UTF-8',xml_declaration=True,standalone=True)
        with ZipFile(target,'w',ZIP_DEFLATED) as zout:
            for item in zin.infolist():zout.writestr(item,edited if item.filename=='word/document.xml' else zin.read(item.filename))
        with ZipFile(target) as zout:
            same=[name for name in zin.namelist() if name!='word/document.xml' and zin.read(name)==zout.read(name)]
            assert len(same)==len(zin.namelist())-1
    write(ROOT/'reviews/DESCRIPTOR_PRESERVATION.json',dict(source=str(source),target=str(target),source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),changed_part='word/document.xml',changed_body_child=12,unchanged_other_parts=len(same),unchanged_other_body_children=len(before)-1))

def instructor_guide(weeks):
    d=setup_doc('Generative AI instructor guide','Twelve week delivery and evidence plan')
    d.add_paragraph('This pack implements a broad and practical Generative AI module for MSc students while preserving the professor’s original descriptor outside section 7.3. Teach the mechanisms with small examples, use current tools for authentic tasks and ask students to explain the evidence behind each result.')
    d.add_heading('What controls the curriculum',1)
    d.add_paragraph('The authoritative template is context/Descriptor - Generative AI-old.docx. The revised copy changes only the weekly content table in section 7.3. The learning outcomes, hours, group project at 60 percent and examination at 40 percent remain unchanged. Sample assessment references and reading lists outside 7.3 are preserved exactly.')
    d.add_heading('Teaching sequence',1)
    for w in weeks:
        d.add_heading(f'Week {w["n"]:02d} {w["title"]}',2)
        d.add_paragraph(w['overview']);d.add_paragraph('Evidence: '+w['lab']['deliverable'])
    d.add_heading('Before teaching',1)
    for text in ['Run each notebook from a fresh kernel. The ordinary offline labs use the Python standard library. The required model-training notebooks need PyTorch in a separate environment.','Pre-run the required VAE and GAN exercise in week 2 and LoRA exercise in week 8. Students must identify actual training, frozen parameters and held-out evidence. Small teaching models do not establish foundation-model performance.','Choose an institution-approved chat or multimodal tool for live tasks. Use fictional inputs. Check availability and costs at the time of teaching and keep an explicitly labelled saved example for outages.','Open the slides in presentation mode and use speaker notes for the explanation and answers. Questions precede answers where possible.','Use the official assessment brief for summative marking. The pack supplies formative practice and an evidence map, without changing assessment requirements.']:
        d.add_paragraph(text,'List Bullet')
    d.add_heading('Learning outcome evidence',1)
    for line in ['LO1: compare VAE, GAN, diffusion and transformer mechanisms, interpret small losses and attention calculations.','LO2: inspect alignment and reasoning across text, image, audio and video through week 3 and week 6 tasks. Retain actual tool outputs when available.','LO3: implement the small generative models, complete required parameter-efficient fine-tuning and explain before-and-after evidence.','LO4: use quantitative checks and a human rubric throughout, then consolidate bias and hallucination evaluation in week 10.','LO5: implement source-grounded retrieval or a bounded tool workflow and deploy the small trained generator through a local interface.']:
        d.add_paragraph(line,'List Bullet')
    d.add_heading('Review loop',1)
    d.add_paragraph('Read a lesson against its mapped outcomes, check each claim and reference, execute the notebook, build the files, inspect the rendered pages and slides, then record and correct issues. Run verification again after changes. Passing code checks does not prove students understand the lesson; use their exit answers and lab explanations to decide the next teaching adjustment.')
    d.add_heading('Scope and companion modules',1)
    d.add_paragraph('AI Technologies already introduces general AI, LLMs and prompting. Begin with a short diagnostic recap and then apply those ideas specifically to generative tasks. Explainable and Emerging AI Technologies covers detailed assurance, regulation and explanation. Generative AI for Business covers organisational adoption and strategy. Here the focus is model understanding, practical generative work and evidence about outputs. MCP and multi-agent architectures are awareness topics, not separate engineering courses.')
    d.save(STAGE/'00_Instructor_Guide.docx')

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--weeks',default='all');ap.add_argument('--documents-only',action='store_true');args=ap.parse_args()
    selected=list(range(1,13)) if args.weeks=='all' else [int(x) for x in args.weeks.split(',')]
    if not selected or any(n<1 or n>12 for n in selected):raise ValueError('Weeks must be 1 to 12')
    STAGE.mkdir(parents=True,exist_ok=True);BUILD.mkdir(parents=True,exist_ok=True)
    allweeks=[read(CANON/f'week_{n:02d}.json') for n in range(1,13)]
    for w in allweeks:
        if w['n'] not in selected:continue
        assert len(w['concepts'])==4 and len(w['practice'])==3
        folder=STAGE/f'week_{w["n"]:02d}';folder.mkdir(exist_ok=True)
        make_notes(w,folder);write(folder/'slides_source.json',make_slides(w))
        lab=CANON/f'week_{w["n"]:02d}_lab.ipynb'
        if lab.exists():shutil.copy2(lab,folder/f'Week_{w["n"]:02d}_Lab.ipynb')
    patch_descriptor();instructor_guide(allweeks)
    if not args.documents_only:
        runtime=Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node.exe'
        if not runtime.exists():raise RuntimeError('Set up the artifact-tool runtime before building PPTX')
        subprocess.run([str(runtime),str(ROOT/'build_slides.mjs'),'--weeks',','.join(map(str,selected))],check=True,cwd=ROOT)
    print('Built staging files at',STAGE)

if __name__=='__main__':main()
