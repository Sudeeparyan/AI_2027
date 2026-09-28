"""Native, editable visual templates for the 40-slide beginner lesson.

Every visual is constrained to the right-hand diagram viewport. course.py exports
that viewport to versioned PNG/SVG assets and embeds the PNG in the classroom copy.
"""
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from . import slides as S
INK=S.INK;BLUE=S.BLUE;GREEN=S.GREEN;ORANGE=S.ORANGE;MUTED=S.MUTED
PALE=S.PALE;LIGHT=S.LIGHT;SOFT=S.SOFT;WHITE=S.WHITE;RED=S.RED
COTTON=RGBColor(255,249,239);ROSE=RGBColor(253,239,237)

def shape(s,x,y,w,h,fill,rounded=False):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=fill;sh.line.fill.background()
    return sh

def txt(s,x,y,w,h,value,size=17,color=INK,bold=False,align=None):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame
    tf.word_wrap=True;tf.margin_left=tf.margin_right=Inches(.025);tf.margin_top=Inches(.01)
    for i,line in enumerate(str(value).split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line
        p.font.name='Aptos';p.font.size=Pt(size);p.font.color.rgb=color;p.font.bold=bold
        if align:p.alignment=align
    return sh

def card(s,x,y,w,h,value,fill=PALE,size=16,color=INK):
    shape(s,x,y,w,h,fill,True);t=txt(s,x+.15,y+.12,w-.3,h-.22,value,size,color,True,PP_ALIGN.CENTER)
    t.text_frame.vertical_anchor=MSO_ANCHOR.MIDDLE
    return t

def line(s,x1,y1,x2,y2,color=BLUE):
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    c.line.color.rgb=color;c.line.width=Pt(2.2)
    tip=s.shapes.add_shape(MSO_SHAPE.CHEVRON,Inches(x2-.09),Inches(y2-.07),Inches(.16),Inches(.14))
    tip.fill.solid();tip.fill.fore_color.rgb=color;tip.line.fill.background()

def shorten(t,n=100):
    t=str(t)
    return t if len(t)<=n else t[:n-1].rstrip()+'…'

def pictogram(s,cx,cy,word,color):
    """Small native PowerPoint pictograms chosen from the analogy's actual nouns."""
    key=str(word).lower()
    if any(t in key for t in ('book','policy','source','page','card','rule','receipt','edition','sheet','form','document','label')):kind='page'
    elif any(t in key for t in ('question','request','text','sentence','prompt','answer','reply','note','brief','message','claim','ticket')):kind='speech'
    elif any(t in key for t in ('image','photo','poster','pixel','icon','figure','draft','picture','sample')):kind='picture'
    elif any(t in key for t in ('model','network','weights','training','latent','embedding','vector','decoder','generator','discriminator','token','attention','mask','code','adapter')):kind='chip'
    elif any(t in key for t in ('check','validation','review','test','decision','approve','flag','gate','exam','feedback','judge')):kind='check'
    elif any(t in key for t in ('number','score','cost','amount','sum','loss','curve','price','calculation','practice','measurement')):kind='chart'
    elif any(t in key for t in ('role','permission','access','staff','student','cabinet','shelf')):kind='lock'
    elif any(t in key for t in ('time','date','day','version','future','calendar')):kind='calendar'
    else:kind='node'
    shape(s,cx-.42,cy-.42,.84,.84,color,True)
    W=WHITE
    if kind=='page':
        shape(s,cx-.2,cy-.25,.4,.52,W,False)
        for k in range(3):shape(s,cx-.14,cy-.16+k*.11,.28,.025,color)
    elif kind=='speech':
        shape(s,cx-.27,cy-.20,.54,.38,W,True)
        for k in range(3):shape(s,cx-.17+k*.13,cy-.02,.045,.045,color,True)
    elif kind=='picture':
        shape(s,cx-.27,cy-.22,.54,.44,W,False)
        sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(cx+.08),Inches(cy-.14),Inches(.1),Inches(.1));sh.fill.solid();sh.fill.fore_color.rgb=color;sh.line.fill.background()
        line(s,cx-.2,cy+.12,cx-.02,cy-.04,color);line(s,cx-.02,cy-.04,cx+.17,cy+.14,color)
    elif kind=='chip':
        shape(s,cx-.22,cy-.22,.44,.44,W,True)
        for dx,dy in [(-.10,-.08),(.08,-.08),(-.10,.10),(.08,.10)]:shape(s,cx+dx-.025,cy+dy-.025,.05,.05,color,True)
    elif kind=='check':
        sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(cx-.25),Inches(cy-.25),Inches(.5),Inches(.5));sh.fill.solid();sh.fill.fore_color.rgb=W;sh.line.fill.background()
        line(s,cx-.14,cy,cx-.02,cy+.10,color);line(s,cx-.02,cy+.10,cx+.18,cy-.13,color)
    elif kind=='chart':
        for k,h in enumerate((.16,.3,.45)):shape(s,cx-.25+k*.18,cy+.24-h,.11,h,W)
    elif kind=='lock':
        shape(s,cx-.2,cy-.02,.4,.26,W,True)
        sh=s.shapes.add_shape(MSO_SHAPE.ARC,Inches(cx-.17),Inches(cy-.28),Inches(.34),Inches(.36));sh.line.color.rgb=W;sh.line.width=Pt(3);sh.fill.background()
    elif kind=='calendar':
        shape(s,cx-.24,cy-.22,.48,.46,W,False)
        shape(s,cx-.24,cy-.13,.48,.07,color)
        for k in range(3):shape(s,cx-.16+k*.13,cy+.03,.05,.05,color,True)
    else:
        for dx,dy in ((-.17,-.08),(.17,-.08),(0,.18)):
            sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(cx+dx-.07),Inches(cy+dy-.07),Inches(.14),Inches(.14));sh.fill.solid();sh.fill.fore_color.rgb=W;sh.line.fill.background()

def draw(s,d,week):
    kind=d['type'];items=d.get('items',[])
    shape(s,5.14,1.47,7.43,5.17,WHITE,True)
    # Visuals use spacing and text labels to explain relationships, never anonymous decoration.
    if kind=='story':
        S.signature(s,week)
    elif kind=='quiz':
        txt(s,5.55,1.73,6.7,.45,'VOTE FIRST  •  GIVE A REASON',16,BLUE,True)
        for i,val in enumerate(items[:2]):
            y=2.42+i*1.58;card(s,5.55,y,6.65,1.22,f'{chr(65+i)}  {shorten(val,115)}',[PALE,COTTON][i],16)
        txt(s,5.6,5.8,6.35,.45,'Which source or test could settle this?',16,GREEN,True)
    elif kind=='steps':
        count=min(len(items),5);gap=.14;hh=min(.9,(4.55-(count-1)*gap)/count)
        for i,val in enumerate(items[:count]):
            y=1.75+i*(hh+gap)
            card(s,5.55,y,.72,hh,str(i+1),[LIGHT,PALE][i%2],18,BLUE)
            card(s,6.47,y,5.62,hh,shorten(val,95),[PALE,LIGHT][i%2],16)
    elif kind=='analogy':
        txt(s,5.55,1.74,6.5,.82,shorten(d.get('caption',''),130),18,BLUE,True)
        for i,val in enumerate(items[:3]):
            x=5.47+i*2.34
            # Three labelled familiar objects form a scene; the next slide maps them precisely.
            pictogram(s,x+1.0,3.20,val,[BLUE,GREEN,ORANGE][i])
            card(s,x,3.78,2.05,1.16,shorten(val,50),[PALE,LIGHT,COTTON][i],16)
            if i<2:line(s,x+2.08,4.35,x+2.31,4.35)
        txt(s,5.55,5.45,6.7,.47,'What matches the technical process?',16,GREEN,True)
    elif kind=='map':
        txt(s,5.55,1.72,6.6,.42,'FAMILIAR SCENE                  TECHNICAL IDEA',15,BLUE,True)
        for i,(left,right) in enumerate(d['pairs'][:3]):
            y=2.28+i*1.15;card(s,5.5,y,2.82,.85,shorten(left,58),PALE,15)
            line(s,8.42,y+.42,8.7,y+.42)
            card(s,8.86,y,3.34,.85,shorten(right,65),LIGHT,15)
        txt(s,5.6,5.89,6.5,.45,'The analogy has a limit. Read it on the left.',15,ORANGE,True)
    elif kind=='flow':
        count=len(items)
        txt(s,5.55,1.72,6.6,.42,'FOLLOW THE DATA THROUGH EACH ARROW',15,BLUE,True)
        if count<=4:
            for i,val in enumerate(items):
                y=2.22+i*.95;card(s,5.75,y,6.05,.66,shorten(val,92),[PALE,LIGHT][i%2],15)
                if i<count-1:line(s,8.77,y+.68,8.77,y+.91)
        else:
            for i,val in enumerate(items[:6]):
                y=2.20+i*.72;card(s,5.72,y,6.08,.52,shorten(val,80),[PALE,LIGHT][i%2],14)
                if i<count-1:line(s,8.75,y+.53,8.75,y+.68)
    elif kind=='work':
        heads=['GIVEN','QUESTION','SUPPORTED ANSWER']
        for i,val in enumerate(items[:3]):
            y=1.88+i*1.48
            txt(s,5.58,y,6.3,.31,heads[i],13,[BLUE,ORANGE,GREEN][i],True)
            card(s,5.54,y+.36,6.6,1.04,shorten(val,145),[PALE,COTTON,LIGHT][i],15)
    elif kind=='reveal':
        txt(s,5.55,1.73,6.4,.4,'TEST THE CLAIM AGAINST THE EVIDENCE',15,BLUE,True)
        txt(s,5.56,2.26,6.2,.35,'TEMPTING CLAIM',13,RED,True)
        card(s,5.52,2.67,6.7,1.11,shorten(items[0],135),ROSE,15)
        txt(s,5.56,4.01,6.2,.35,'SUPPORTED EXPLANATION',13,GREEN,True)
        card(s,5.52,4.43,6.7,1.10,shorten(items[1],135),LIGHT,15)
        txt(s,5.6,5.82,6.45,.62,shorten(d.get('caption',''),100),13,MUTED)
    elif kind=='compare':
        card(s,5.45,2.28,3.08,2.25,d['left'],PALE,18)
        card(s,9.02,2.28,3.08,2.25,d['right'],LIGHT,18)
        txt(s,5.52,5.15,6.55,.67,d.get('caption',''),16,GREEN,True)
    elif kind=='table':
        rows=d['rows'];n=len(rows[0]);cw=6.5/n
        for i,row in enumerate(rows[:5]):
            for j,val in enumerate(row):card(s,5.5+j*cw,1.96+i*.84,cw-.08,.7,shorten(val,28),PALE if i==0 else (LIGHT if i%2 else SOFT),14)
    else:raise ValueError(kind)

def create_deck(w,out,records):
    p=Presentation();p.slide_width=Inches(13.333);p.slide_height=Inches(7.5)
    n=w['n'];assert len(records)==40
    stage_color={'open':BLUE,'question':ORANGE,'analogy':ORANGE,'mapping':GREEN,'mechanism':BLUE,'worked':GREEN,'reveal':GREEN,'practice':BLUE,'check':ORANGE}
    for i,record in enumerate(records.values(),1):
        s=p.slides.add_slide(p.slide_layouts[6]);stage=record['teaching_stage'];accent=stage_color[stage]
        shape(s,0,0,13.333,.12,accent)
        txt(s,.65,.34,11.9,.63,record['title'],27,INK,True)
        shape(s,.65,1.49,.07,4.97,accent)
        txt(s,.92,1.5,3.84,.35,stage.upper()+'  /  WEEK '+str(n).zfill(2),11,accent,True)
        body=record['explanation'];size=18 if len(body)<250 else 16.7
        txt(s,.92,2.02,3.89,4.36,body,size,INK)
        shape(s,.65,6.99,11.92,.018,PALE)
        txt(s,.68,7.12,11.84,.24,f'GENERATIVE AI 2027    •    WEEK {n:02d}    •    {i:02d} / 40',9,MUTED)
        draw(s,record['diagram'],n)
        s.notes_slide.notes_text_frame.text=record['speaker_notes']
    out.mkdir(parents=True,exist_ok=True);p.save(out/f'Week_{n:02d}_Slides.pptx')
