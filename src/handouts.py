from docx import Document
from docx.shared import Inches as DI, Pt as DP, RGBColor as DC
SOURCES={}
def make_doc(w,dir):
 d=Document();sec=d.sections[0];sec.top_margin=DI(.72);sec.bottom_margin=DI(.68);sec.left_margin=DI(.86);sec.right_margin=DI(.86)
 styles=d.styles;styles['Normal'].font.name='Aptos';styles['Normal'].font.size=DP(10)
 styles['Heading 1'].font.name='Aptos Display';styles['Heading 1'].font.size=DP(17);styles['Heading 1'].font.color.rgb=DC(31,95,142)
 styles['Heading 2'].font.size=DP(12);styles['Heading 2'].font.color.rgb=DC(20,132,130)
 d.add_heading(f'Week {w["n"]:02d}: {w["title"]}',0)
 d.add_paragraph('Generative AI | MSc teaching notes and further reading | 2027')
 d.add_heading('Learning target',1);d.add_paragraph(w['goal'])
 d.add_heading('Context and teaching plan',1)
 d.add_paragraph(w['case'])
 t=d.add_table(rows=1, cols=2);t.style='Light Shading Accent 1';t.rows[0].cells[0].text='Contact time';t.rows[0].cells[1].text='Suggested use'
 for a,b in [('Two hours: concepts','Case introduction, four ideas, worked example and misconception.'),('Two hours: practice','10 min recall, 20 min demo, 40 min guided practice, 25 min independent change, 15 min review, 10 min exit check.')]:r=t.add_row().cells;r[0].text=a;r[1].text=b
 d.add_heading('Explain the key ideas',1)
 for i,(term,definition) in enumerate(w['terms']):
  d.add_heading(term,2);d.add_paragraph(definition)
  d.add_paragraph(f'Teacher question: Where does {term.lower()} appear in the weekly case, and what could go wrong if we omit it?')
 d.add_heading('Worked example and answer guide',1)
 for label,s in zip(['Input','Change or question','Supported observation'],w['example']):d.add_paragraph(f'{label}: {s}',style='List Bullet')
 d.add_paragraph('Reasoning: state the given evidence first, then the calculation or rule, then the conclusion. If a value is absent, say what cannot be concluded. Keep any simulated mechanism distinct from a trained model or deployed service.')
 if w.get('design'):
  de=w['design'];d.add_heading('Teach the week as a story',0)
  d.add_paragraph(de['story']+': '+de['hook'])
  d.add_paragraph('Ask for a quick prediction before each technical explanation. Return to the prediction after the example and ask what evidence changed the answer. The analogies below are memory aids, and each has a stated limit.')
  for (term,definition),c in zip(w['terms'],de['concepts']):
   d.add_heading(term+': analogy, mechanism and check',1)
   d.add_paragraph('Familiar picture: '+c['analogy'])
   table=d.add_table(rows=1,cols=2);table.style='Light Shading Accent 1';table.rows[0].cells[0].text='Familiar part';table.rows[0].cells[1].text='Technical meaning'
   for a,b in c['mapping']:
    row=table.add_row().cells;row[0].text=a;row[1].text=b
   d.add_paragraph('Where this analogy breaks: '+c['limit'])
   d.add_paragraph('Actual idea: '+definition)
   d.add_paragraph('Work it aloud: '+c['worked'])
   d.add_paragraph('Ask before revealing: '+c['question'])
   d.add_paragraph('Supported response: '+c['answer'])
   d.add_paragraph('Look for this misconception: '+c['wrong'])
  d.add_heading('New situation for discussion',1);d.add_paragraph(de['transfer'])
  d.add_paragraph('Memory cue: '+de['memory'])
 d.add_heading('Guided laboratory',1);d.add_paragraph(w['lab'])
 for i,s in enumerate(w['steps'],1):d.add_paragraph(f'{i}. {s}')
 d.add_paragraph(f'Individual change: {w["change"]} Expected evidence: {w["output"]}')
 d.add_heading('Checks and common error',1)
 d.add_paragraph(f'Measure: {w["metric"]}. Record the input, setting, source version, result and the relevant numerator and denominator.')
 d.add_paragraph(f'Plausible incorrect claim: {w["misconception"]}')
 d.add_paragraph(f'Correction: {w["correction"]}')
 d.add_paragraph('If learners disagree, reproduce the same case and compare the underlying source or expected answer. Record unresolved cases for instructor review.')
 d.add_heading('Further discussion',1)
 for q in [f'How could the conclusion change if the source or input changes?',f'Which part of the result is observed and which part is an assumption?',f'What simple test would detect the failure described above?']:
  d.add_paragraph(q,style='List Bullet')
 d.add_heading('References and optional extension',1)
 d.add_paragraph('Core readings are official vendor documentation or original research. Read the relevant section; an account or paid API is unnecessary for this notebook. Page details may change before 2027.')
 for k in w['reading']:
  d.add_paragraph(f'{k}. {SOURCES[k]}')
 d.add_paragraph('Optional extension: replace one classroom proxy with a locally approved model or service, keep the original baseline, repeat the same checks, and state the additional privacy and cost assumptions.')
 d.add_paragraph('Curriculum scope: the supplied 12-week teaching sequence, recorded in docs/SYLLABUS_ALIGNMENT.md. The original module descriptor is retained in sources/originals/07-Descriptor-Generative-AI-new-.docx. Supporting project notes: Deep-reserach.txt, RAG.txt, AI-Agents.txt, MCP-&-Vectore-Database.txt, Muti-Agent-AI-System.txt, Eval-&-Finetune&-Loop-Enginnering.txt. These are planning sources, not empirical performance evidence.')
 if w.get('study'):
  st=w['study'];d.add_heading('Student revision and exam preparation',0)
  d.add_paragraph('Outcomes: '+st['outcomes']+' | 2 hours concepts + 2 hours guided practice')
  d.add_heading('The idea in plain language',1);d.add_paragraph(st['plain_explanation'])
  d.add_heading('A worked example',1);d.add_paragraph(st['worked_example'])
  d.add_heading('Practice before reading the answers',1)
  for i,q in enumerate(st['exam'],1):d.add_paragraph(f'{i}. {q["question"]} ({q["marks"]} suggested practice marks)')
  d.add_paragraph(st['study_task'])
  d.add_heading('Model answers and marking guidance',1)
  for i,q in enumerate(st['exam'],1):d.add_paragraph(f'{i}. {q["model_answer"]}')
  d.add_paragraph('For each four-mark practice answer: two marks for the correct idea, one for a relevant example or calculation, one for an appropriate explanation or limitation. These suggested practice marks do not change the module assessment scheme.')
  d.add_heading('Official reading and video support',1)
  for ref in st['resources']:d.add_paragraph(f'{ref["id"]}  {ref["publisher"]}: {ref["title"]}\n{ref["url"]}\nStudy focus: {ref["use"]}')
  d.add_paragraph(st['reference_note'])
 for para in d.paragraphs:para.paragraph_format.keep_together=True
 for border in d.styles.element.xpath('.//w:pBdr'):
  border.getparent().remove(border)
 d.save(dir/f'Week_{w["n"]:02d}_Reading_and_Teaching_Notes.docx')
