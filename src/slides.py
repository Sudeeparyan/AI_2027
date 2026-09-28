"""Editable visual PowerPoint lessons for the 12-week Generative AI module.

Usage:
    python course.py improve --weeks all

The module creates PowerPoint-native diagrams and presenter scripts. Edit
the weekly JSON files in content/, then rerun.
"""
from pathlib import Path
import textwrap
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.dml.color import RGBColor

INK=RGBColor(24,43,64); MUTED=RGBColor(68,86,101)
BLUE=RGBColor(37,94,156); GREEN=RGBColor(17,126,110)
ORANGE=RGBColor(199,104,44); PALE=RGBColor(235,245,251)
LIGHT=RGBColor(239,247,243); WHITE=RGBColor(255,255,255)
SOFT=RGBColor(248,250,252); RED=RGBColor(183,73,69)

# Exact labels in these diagrams carry the week-specific technical content.
# No deck relies on unlabeled decoration to explain an idea.
VISUALS={}
OVERRIDES={}
ASSET_ROOT=None

def box(slide,x,y,w,h,label,fill=PALE,sz=17):
 s=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
 s.fill.solid();s.fill.fore_color.rgb=fill;s.line.color.rgb=fill
 tf=s.text_frame;tf.clear();tf.word_wrap=True;tf.vertical_anchor=MSO_ANCHOR.MIDDLE
 tf.margin_left=tf.margin_right=Inches(.10)
 p=tf.paragraphs[0];p.text=label;p.alignment=PP_ALIGN.CENTER
 p.font.name='Aptos';p.font.size=Pt(sz);p.font.bold=True;p.font.color.rgb=INK
 return s
def label(slide,x,y,w,h,text,size=19,color=INK,bold=False):
 s=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
 t=s.text_frame;t.word_wrap=True;t.margin_left=t.margin_right=0
 for j,line in enumerate(str(text).split('\n')):
  p=t.paragraphs[0] if j==0 else t.add_paragraph()
  p.text=line;p.font.name='Aptos';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color
 return s
def arrow(slide,x1,y1,x2,y2,color=BLUE):
 c=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
 c.line.color.rgb=color;c.line.width=Pt(2.2)
 e=slide.shapes.add_shape(MSO_SHAPE.CHEVRON,Inches(x2-.09),Inches(y2-.075),Inches(.17),Inches(.15))
 e.fill.solid();e.fill.fore_color.rgb=color;e.line.fill.background()
def chain(slide,items,y=2.3,colors=None):
 items=[str(i) for i in items];n=len(items);gap=.21
 bw=(7.25-(n-1)*gap)/n
 for j,item in enumerate(items):
  x=5.25+j*(bw+gap)
  box(slide,x,y,bw,1.22,item,(colors or [PALE,LIGHT])[j%2],13 if n>=5 else 15)
  if j:arrow(slide,x-gap+.02,y+.61,x-.035,y+.61)
def split(slide,left,right,foot=''):
 box(slide,5.3,2.0,3.35,2.25,left,PALE,18)
 box(slide,9.05,2.0,3.35,2.25,right,LIGHT,18)
 if foot:label(slide,5.35,4.65,7.1,.75,foot,17,GREEN,True)
def branch(slide,source,yes,no):
 box(slide,7.1,1.75,3.1,1.05,source,PALE,18)
 arrow(slide,8.1,2.82,6.65,3.53);arrow(slide,9.2,2.82,10.65,3.53)
 box(slide,5.2,3.55,3.0,1.27,yes,LIGHT,16)
 box(slide,9.15,3.55,3.0,1.27,no,RGBColor(251,238,235),16)
def stacked(slide,items):
 for j,it in enumerate(items[:5]):
  y=1.72+j*.77
  box(slide,5.35,y,1.0,.56,str(j+1),LIGHT,17)
  label(slide,6.55,y-.02,5.5,.58,it,17)
def score(slide,values,caption):
 for j,(name,num,den) in enumerate(values):
  y=1.75+j*1.05;label(slide,5.25,y,2.2,.45,name,17)
  box(slide,7.55,y,3.5*num/max(den,1),.42,'',GREEN,10)
  label(slide,11.25,y,1.1,.45,f'{num}/{den}',17,INK,True)
 label(slide,5.35,5.6,6.9,.66,caption,17,MUTED)

def signature(slide,n):
 """A topic-specific teaching picture for the opening slide of each week."""
 if n==1 and ASSET_ROOT and (ASSET_ROOT/'illustrations'/'student_evidence_v1.png').exists():
  slide.shapes.add_picture(str(ASSET_ROOT/'illustrations'/'student_evidence_v1.png'),Inches(5.25),Inches(1.7),width=Inches(7.15),height=Inches(4.65))
 elif n==1:
  box(slide,5.3,1.9,2.45,1.1,'Student question',PALE)
  box(slide,9.45,1.9,2.6,1.1,'Candidate answer',LIGHT)
  arrow(slide,7.8,2.45,9.35,2.45)
  box(slide,7.3,4.0,2.65,1.05,'Policy P1',PALE)
  arrow(slide,8.55,4.0,10.55,3.15)
  label(slide,6.1,5.6,5.5,.5,'Check the answer against the policy',18,GREEN,True)
 elif n==2:
  label(slide,5.3,1.65,7,.5,'VAE: reconstruct an example',18,BLUE,True)
  for x,t in zip([5.3,7.7,10.1],['Image','Encoder + z','Decoder → image']):box(slide,x,2.25,2.1,1.0,t,PALE,15)
  arrow(slide,7.43,2.75,7.63,2.75);arrow(slide,9.83,2.75,10.03,2.75)
  label(slide,5.3,3.65,7,.5,'GAN: improve a proposed example',18,GREEN,True)
  for x,t in zip([5.3,7.7,10.1],['Random z','Generator','Discriminator']):box(slide,x,4.2,2.1,1.0,t,LIGHT,15)
  arrow(slide,7.43,4.7,7.63,4.7);arrow(slide,9.83,4.7,10.03,4.7)
  label(slide,7.75,5.5,4.8,.45,'Feedback updates both networks',15,GREEN,True)
 elif n==3:
  for group,noise in enumerate([0,1,2]):
   x0=5.55+group*2.3
   for r in range(4):
    for c in range(4):
     bright=(r==2 or c==1)
     base=(215 if bright else 48)
     random_pixel=25+((r*97+c*61+r*c*39+group*53)%210)
     strength=[0,.42,.85][group]
     val=round((1-strength)*base+strength*random_pixel)
     sq=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x0+c*.32),Inches(2.25+r*.32),Inches(.29),Inches(.29))
     sq.fill.solid();sq.fill.fore_color.rgb=RGBColor(val,val,val);sq.line.fill.background()
   label(slide,x0,3.75,1.65,.7,['clean','less noise','more noise'][group],15,INK)
  label(slide,5.45,5.1,6.9,.65,'Generation uses learned reverse steps',19,GREEN,True)
 elif n==4:
  label(slide,5.25,1.6,6.6,.55,'Which earlier token can each word use?',18,BLUE,True)
  vals=[[1,0,0],[.3,.7,0],[.25,.65,.1]]
  for c,t in enumerate(['I','read','books']):label(slide,6.15+c*1.14,2.05,1.02,.38,t,13,MUTED,True)
  for r in range(3):
   label(slide,5.22,2.58+r*.79,.8,.35,['I','read','books'][r],13,MUTED,True)
   for c in range(3):
    v=vals[r][c]
    shade=int(246-v*110)
    box(slide,6.15+c*1.14,2.42+r*.79,1.02,.67,f'{v:.2f}',RGBColor(shade,242,240),15)
  label(slide,5.45,5.4,7.1,.65,'The mask hides future positions during generation',17,GREEN,True)
 elif n in (5,10):
  box(slide,5.35,1.65,6.85,3.9,'',PALE)
  label(slide,5.75,1.94,6,.55,'INVOICE  /  INV-04',20,BLUE,True)
  for i,(a,b) in enumerate([('Quantity','2'),('Unit price','€20'),('Tax','€5'),('Listed total','€54')]):
   y=2.65+i*.57;label(slide,5.9,y,3.3,.4,a,16);label(slide,9.7,y,1.55,.4,b,17,RED if i==3 else INK,True)
  label(slide,5.55,5.75,6.3,.5,'Calculated total = €45. Flag the mismatch.',17,GREEN,True)
 elif n in (6,7,12):
  for i,t in enumerate(['P1 Travel','P2 Library','P3 Deadline','P4 Appeal','P5 Staff']):
   box(slide,5.32+i*1.45,2.05,1.28,1.6,t,PALE if i<4 else RGBColor(251,238,235),13)
  label(slide,5.35,4.2,7.0,.55,'Retrieve only the policy allowed for this user',17,GREEN,True)
  chain(slide,['Question','Allowed policy','Cited answer'],5.12)
 elif n==8:
  branch(slide,'Action requested','Read-only: calculate','Write: seek approval')
  label(slide,5.45,5.2,6.7,.58,'The model cannot grant itself permission',18,GREEN,True)
 elif n==9:
  score(slide,[('Base',6,10),('Prompt',8,10),('Tuned',9,10)],'Practice scores only: test on fresh cases before choosing.')
 elif n==11:
  for j,(t,clr) in enumerate([('Version 2',PALE),('Access test fails',RGBColor(251,238,235)),('Keep version 1',LIGHT)]):
   box(slide,5.38+j*2.36,2.45,2.07,1.42,t,clr,16)
   if j:arrow(slide,5.38+j*2.36-.19,3.16,5.38+j*2.36-.04,3.16)
  label(slide,5.4,5.3,6.7,.58,'A failed permission test blocks release',18,GREEN,True)
 else:chain(slide,['Input','Check','Output'],2.75)

def new_slide(p,title,lead,n,i,visual,notes):
 override=OVERRIDES.get(f's{i:02d}',{})
 title=override.get('title',title);lead=override.get('explanation',lead);notes=override.get('speaker_notes',notes)
 s=p.slides.add_slide(p.slide_layouts[6])
 label(s,.7,.35,11.7,.7,title,28,BLUE,True)
 box(s,.62,1.46,3.98,4.96,'',SOFT,17)
 label(s,.95,1.84,3.35,4.15,lead,19,INK)
 label(s,.74,7.12,11.8,.18,f'GENERATIVE AI   /   WEEK {n:02d}   /   {i:02d}',9,MUTED)
 if override.get('diagram'): custom_visual(s,override['diagram'])
 else: visual(s)
 s.notes_slide.notes_text_frame.text=notes
 return s

def custom_visual(slide,spec):
 kind=spec['type']
 if kind=='flow':chain(slide,spec['items'],2.65)
 elif kind=='steps':
  for i,item in enumerate(spec['items']):
   box(slide,5.28,1.72+i*1.24,7.12,1.02,item,PALE if i%2==0 else LIGHT,15)
 elif kind=='compare':split(slide,spec['left'],spec['right'],spec.get('caption',''))
 elif kind=='table':
  rows=spec['rows'];cols=len(rows[0]);bw=7.15/cols
  for i,row in enumerate(rows):
   for j,value in enumerate(row):
    box(slide,5.25+j*bw,1.7+i*.77,bw-.07,.66,str(value),PALE if i==0 else LIGHT,14)
  if spec.get('caption'):label(slide,5.3,1.8+len(rows)*.77,7,.7,spec['caption'],15,MUTED)
 elif kind=='loss':
  label(slide,5.3,1.65,7,.5,'Supplied VAE errors (illustrative)',17,BLUE,True)
  x0,y0=6.05,5.0;ww,hh=5.3,2.5
  for x1,y1,x2,y2 in [(x0,y0,x0+ww,y0),(x0,y0,x0,y0-hh)]:
   a=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));a.line.color.rgb=MUTED
  for values,color,name in [([.30,.18,.12],BLUE,'training'),([.32,.20,.26],ORANGE,'validation')]:
   pts=[(x0+i*ww/2,y0-v/.4*hh) for i,v in enumerate(values)]
   for (x1,y1),(x2,y2) in zip(pts,pts[1:]):
    a=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));a.line.color.rgb=color;a.line.width=Pt(2)
   for (x,y),v in zip(pts,values):label(slide,x-.1,y+(.06 if name=='training' else -.4),.8,.4,str(v),13,color,True)
   label(slide,10.35,5.55 if name=='training' else 5.98,2.0,.35,name,14,color,True)
  for i in range(3):label(slide,x0+i*ww/2-.3,5.15,1,.3,f'epoch {i+1}',12)
  label(slide,5.3,2.0,.7,.4,'error',12,MUTED)
 else:raise ValueError('Unknown diagram type: '+kind)

def instruction(w,focus,visual,example,question):
 return (f'SAY: {focus} Start with this week’s college example: {w["case"]} '
         f'Point to the diagram while you trace {visual}. Work the concrete example aloud: {example} '
         f'PAUSE: {question} Ask learners to give a reason based on the visible source or number. '
         f'CHECK: Have one learner explain each arrow without reading the slide. '
         f'If an answer depends on a model, remind students where this classroom notebook uses a simplified stand-in. '
         f'NEXT: Run the corresponding notebook cell, then let learners change one input and record what happened.')

def create_deck(w,out):
 n=w['n'];v=VISUALS[n];p=Presentation();p.slide_width=Inches(13.333);p.slide_height=Inches(7.5)
 def put(title,lead,draw,focus,visual,question):
  ex=' '.join(w['example'])
  return new_slide(p,title,lead,n,len(p.slides)+1,draw,instruction(w,focus,visual,ex,question))
 put(f'Week {n:02d}: {w["title"]}',w['goal'],lambda s:signature(s,n),
     'Today we move through the diagram from a real input to a checked outcome.', 'the complete path from left to right', 'Which step would you check first?')
 put('The question we will solve',w['case'],lambda s:chain(s,['Question','Evidence','Decision'],2.7),
     'Begin with an ordinary human decision before introducing a model.', 'the three blocks as a person would solve the problem', 'What would count as a trustworthy answer?')
 put('What we will learn',w['goal']+'\n\nExplain the diagram, test one example, and justify the result.',lambda s:stacked(s,['Explain the idea','Run the notebook','Change one input','Check the result']),
     'The course rewards an explanation supported by a result.', 'the four learning actions', 'Which action helps us find an error?')
 put('Four ideas in this lesson','Each idea has a specific job. We will trace each one through the same case.',lambda s:chain(s,[x[0] for x in w['terms']],2.8),
     'Read the four short labels before the definitions.', 'the four parts in order', 'Which label is new to you?')
 for j,(term,definition) in enumerate(w['terms']):
  put(f'{term}: a picture',definition,lambda s,j=j:chain(s,v['term'][j],2.75),
      definition, 'each arrow of this specific component diagram', f'What enters and leaves {term.lower()}?')
 put('The complete process','These blocks show where today’s input goes and what must be checked before we trust the result.',lambda s:chain(s,v['path'],2.68),
     'Do not skip the final check simply because the middle blocks produced an output.', 'every stage of the week-specific process', 'Where could a wrong answer first appear?')
 put('Where does the evidence travel?','Follow the original source through the process. Each later claim must still be traceable to that source.',
     lambda s:chain(s,['Source ID',v['path'][1],v['path'][-1],'Evidence record'],2.75),
     'Trace the source identifier even when the intermediate result looks plausible.', 'the evidence path from original source to final record', 'Can you recover the original source from the result?')
 put('Step 1: what enters?',w['example'][0],lambda s:chain(s,['Original input',v['path'][0],'Source or value'],2.8),
     'Locate the exact source or number in the worked example.', 'how the source becomes an input', 'Which part must be copied exactly?')
 put('Step 2: what changes?',w['example'][1],lambda s:chain(s,[v['path'][0],v['path'][1],v['path'][-1]],2.8),
     'Name the one condition that changes while the comparison stays fair.', 'the intermediate processing stage', 'What remains fixed?')
 put('Step 3: what can we conclude?',w['example'][2],lambda s:branch(s,'Check the evidence',v['good'],'More review needed'),
     'Only make a claim supported by the given example.', 'the supported and unsupported paths', 'Which words in the example prove the answer?')
 put('Compare two outcomes','Use one fixed case to see how these outcomes differ.',lambda s:split(s,*v['compare'],'Check each outcome against the same evidence.'),
     'Compare the two labeled outcomes using one fixed input rather than two unrelated examples.', 'the two alternatives and the shared evidence', 'What evidence separates these outcomes?')
 put('A common mistake',w['misconception'],lambda s:branch(s,'Is the claim supported?',v['good'],v['bad']),
     'Present the wrong answer as a genuine hypothesis, not a trick question.', 'why one branch needs evidence and the other is unsafe', 'What counterexample would refute the claim?')
 put('The corrected explanation',w['correction'],lambda s:chain(s,['Observed fact','Check or calculation','Limited conclusion'],2.8),
     w['correction'], 'the difference between an observation and a conclusion', 'Which part is still uncertain?')
 put('Notebook: run the starter',w['lab'],lambda s:stacked(s,['Open the week’s notebook','Run cells from the top','Read the printed output','Compare it with the example']),
     'Demonstrate the first run before allowing individual changes.', 'the actual student workflow', 'Does your output match the worked example?')
 put('Notebook: inspect the code','Find the input, the transformation, the check, and the result. These are the four lines to narrate.',lambda s:chain(s,['Input data','Python function','Validation check','Printed result'],2.75),
     'Point to the actual code cell and identify the variable for the input.', 'the code execution path', 'Which line would you modify for a new input?')
 put('Change just one condition',w['change'],lambda s:split(s,'Before\noriginal input','After\none change','Keep everything else fixed.'),
     'Have learners predict the output before running the cell again.', 'the controlled before and after comparison', 'Did the output change as predicted?')
 put('Record your observation',w['output'],lambda s:chain(s,['Input','Setting','Observed result','Explanation'],2.8),
     'A result without its input and setting cannot be reproduced.', 'the four columns of the evidence record', 'Can your partner reproduce your result?')
 put('Measure the result',w['metric'],lambda s:score(s,[('Example A',3,4),('Example B',2,4)],'Illustrative counts: use the notebook’s real numerator and denominator.'),
     f'Define the real measure as {w["metric"]}; the bars are illustrative, not experimental results.', 'numerators and denominators on the two illustrative bars', 'What does a missed case change?')
 put('Look at a failure','Incorrect claim: '+w['misconception']+'\n\nCorrection: '+w['correction'],lambda s:chain(s,['Unexpected result','Find the failing block','Design a test'],2.75),
     'An error is evidence about a specific layer; avoid saying only that the model failed.', 'the point of failure and a targeted test', 'Which layer would you inspect?')
 put('Check access and sources','Use approved fictional examples. Verify source IDs and user role before showing an answer or using a tool.',lambda s:branch(s,'Permission check','Allowed: continue','Denied: stop'),
     'Even a correct answer is unacceptable if it exposes material to the wrong person.', 'the permission branch before an answer is shown', 'Who is authorized to read this source?')
 put('When a person reviews','A person checks an unsupported claim, a missing field, or an action that would change a record.',lambda s:chain(s,['Flag uncertainty','Person checks evidence','Approve or correct'],2.8),
     'Give one concrete reason for review from today’s case.', 'how a human decision is connected to evidence', 'What would you show the reviewer?')
 put('Try these two questions',f'1. Explain: {w["terms"][0][0]}.\n\n2. Why is this wrong? {w["misconception"]}',lambda s:split(s,'Explain in your own words','Point to evidence','Pair discussion: 2 minutes'),
     'Give students time to answer before revealing a sample explanation.', 'the explanation and supporting-evidence tasks', 'Can you answer without reading the definition?')
 put('What to hand in',w['output'],lambda s:chain(s,['Notebook result','Changed input','Error explanation','Source or test'],2.8),
     'Weekly evidence is short and formative; it builds toward the final project.', 'the four pieces of evidence', 'Is each claim reproducible from your file?')
 put('What to remember',w['goal']+'\n\nThe useful question is: what evidence supports the result?',lambda s:chain(s,[v['path'][0],v['path'][-1],'Verify'],2.9),
     'Recap the journey from question to checked output.', 'the beginning, result, and verification block', 'What would you do if the evidence disappeared?')
 put('Read before next class',f'Read one selected section from {w["reading"][0]}.\n\nBring one question about this week’s example.',lambda s:stacked(s,['Open weekly Word guide','Read the linked source','Bring one question']),
     'Point students to the actual link in the weekly Word guide.', 'the short preparation sequence', 'Which term deserves another example?')
 assert len(p.slides)==28,(n,len(p.slides))
 out.mkdir(parents=True,exist_ok=True);p.save(out/f'Week_{n:02d}_Slides.pptx')
